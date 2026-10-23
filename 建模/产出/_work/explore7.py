# -*- coding: utf-8 -*-
# 本程序及代码是在人工智能工具辅助下完成的。
# AI 工具名称：DeepSeek Harness；模型：deepseek-flash；开发机构/公司：DeepSeek。
"""TM1-01 探查 7：4.csv 字段 / summary 表头 / 8.xlsx 感官名 / 阈值表结构。"""
import pandas as pd
import openpyxl

BASE = "数据/A/"
df4 = pd.read_csv(BASE + "4.csv", encoding="utf-8-sig")
cols4 = [c for c in df4.columns if c not in ("beer", "beer_id", "tasting_category_fine")]
print("4.csv n cols:", len(cols4), " shape:", df4.shape)
for i, c in enumerate(cols4, 1):
    print(f"{i:2d} {c}")

wb3 = openpyxl.load_workbook(BASE + "3.xlsx", data_only=False)
ws = wb3["summary"]
hdr = [ws.cell(1, c).value for c in range(2, 52)]
print("\nsummary headers:", len(hdr))
print(hdr)

wb8 = openpyxl.load_workbook(BASE + "8.xlsx", data_only=True)
ws8 = wb8["Sensory thresholds"]
print("\nthresholds header:", [ws8.cell(1, c).value for c in range(1, ws8.max_column + 1)])
print("rows:", ws8.max_row)
for r in range(2, 6):
    print([ws8.cell(r, c).value for c in range(1, ws8.max_column + 1)])

wsn = wb8["Sensory names"]
print("\nSensory names header:", [wsn.cell(1, c).value for c in range(1, wsn.max_column + 1)])
rnames = [(wsn.cell(r, 1).value, wsn.cell(r, 2).value) for r in range(2, wsn.max_row + 1)]
from collections import Counter
print("datasets:", Counter(d for d, _ in rnames))
trained = [n for d, n in rnames if d == "Trained panel"]
print("trained count:", len(trained))
print(trained)
