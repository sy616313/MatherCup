# -*- coding: utf-8 -*-
# 本程序及代码是在人工智能工具辅助下完成的。
# AI 工具名称：DeepSeek Harness；模型：deepseek-flash；开发机构/公司：DeepSeek。
"""TM1-01 生成 维度映射 与 缺陷阈值表。"""
import re

import openpyxl
import pandas as pd

BASE = "数据/A/"
meta = ("beer", "beer_id", "tasting_category_fine")

df4 = pd.read_csv(BASE + "4.csv", encoding="utf-8-sig")
cols4 = [c for c in df4.columns if c not in meta]

wb3 = openpyxl.load_workbook(BASE + "3.xlsx", data_only=False)
ws3 = wb3["summary"]
sum_cols = []
for c in range(2, 52):
    letter = openpyxl.utils.get_column_letter(c)
    name = ws3.cell(1, c).value
    formula = ws3.cell(2, c).value
    m = re.search(r"\$([A-Z]+)\$(\d+)", formula or "")
    sum_cols.append(dict(name=name, letter=letter, cell=f"{m.group(1)}{m.group(2)}" if m else "",
                         formula=formula, row1=ws3.cell(1, c).value))

wb8 = openpyxl.load_workbook(BASE + "8.xlsx", data_only=True)
wsn = wb8["Sensory names"]
names = {}
for r in range(2, wsn.max_row + 1):
    ds, rname = wsn.cell(r, 1).value, wsn.cell(r, 2).value
    if ds == "Trained panel":
        names[rname] = dict(paper=wsn.cell(r, 5).value, long=wsn.cell(r, 6).value,
                            panel=wsn.cell(r, 7).value)

SPELL = {"A_esters_fruity": "A_ester_fruity", "F_esters_fruity": "F_ester_fruity", "X4vg": "4vg"}
by_name = {s["name"]: s for s in sum_cols}
spell_cells = [SPELL[c] for c in cols4 if c in SPELL]
c_group_start = cols4.index("acidity") + 1

rows = []
for i, c4 in enumerate(cols4, 1):
    key = SPELL.get(c4, c4)
    s = by_name.get(key)
    note = []
    if c4 in SPELL:
        note.append(f"拼写差异：summary 写作 `{s['name']}`")
    if i >= c_group_start:
        note.append("C 组：summary 列序与 4.csv 不同，禁止按位置取")
    nm = names.get(c4, {})
    cell = s["cell"]
    if c4 == "barnyard":
        cell = "AG90"
        note.append("⚠ summary 实际错引 `AG76`（与 stale_hops 同格），**应为 `AG90`**，见 §三")
    rows.append(dict(序号=i, c4=c4, summary=s["name"], letter=s["letter"], cell=cell,
                     readable=nm.get("paper", ""), note="；".join(note)))

pd.DataFrame(rows).to_csv("建模/产出/_work/map50.csv", index=False, encoding="utf-8", lineterminator="\n")
print("rows:", len(rows), " unmatched:", [r["c4"] for r in rows if not r["summary"]])
print("spelling diffs:", [r["c4"] for r in rows if r["c4"] in SPELL])
print("barnyard row:", [r for r in rows if r["c4"] == "barnyard"])
print("stale row:", [r for r in rows if r["c4"] == "stale_hops"])

# ---------- 阈值表 ----------
ws = wb8["Sensory thresholds"]
hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
idx = {h: i + 1 for i, h in enumerate(hdr)}
keep = ["acetaldehyde", "2,3-butanedione", "DES", "DMS", "H2S", "SO2-mg.L.", "X4VG"]
out = []
for r in range(2, ws.max_row + 1):
    py = ws.cell(r, idx["field_name_python"]).value
    if py in keep:
        out.append(dict(化合物=ws.cell(r, idx["chemical_name"]).value,
                        field_name_python=py,
                        field_name_R=ws.cell(r, idx["field_name_R"]).value,
                        数值=ws.cell(r, idx["threshold_mean"]).value,
                        中位数=ws.cell(r, idx["threshold_median"]).value,
                        单位=ws.cell(r, idx["unit"]).value,
                        出处=f"8.xlsx/Sensory thresholds 第 {r} 行"))
t = pd.DataFrame(out)
print("\nthreshold rows kept:", len(t))
print(t.to_string(index=False))
t.to_csv("建模/产出/_work/thr.csv", index=False, encoding="utf-8", lineterminator="\n")
