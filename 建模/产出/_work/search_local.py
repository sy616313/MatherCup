# -*- coding: utf-8 -*-
"""S1 检索 · 第一步：仓库内全部资料（技能 + 精选 PDF + 8.xlsx 其他 sheet）。"""
import pathlib
import re

import openpyxl
import pymupdf

ROOT = pathlib.Path(".")
KEY = re.compile(r"nonenal|methanethiol|MeSH|stale hop|臭|阈值|threshold", re.I)

print("== 1) 技能与文档（.md） ==")
hits = 0
for p in list((ROOT / "资料").rglob("*.md")) + list((ROOT / "建模").rglob("*.md")):
    t = p.read_text(encoding="utf-8", errors="ignore")
    for m in KEY.finditer(t):
        s = max(0, m.start() - 40)
        print(f"  {p}: …{t[s:m.start()+60].replace(chr(10),' ')}…")
        hits += 1
        break
print("  hits:", hits)

print("\n== 2) 8.xlsx 全部 sheet 关键词 ==")
wb = openpyxl.load_workbook("数据/A/8.xlsx", data_only=True)
for sh in wb.sheetnames:
    ws = wb[sh]
    found = 0
    for row in ws.iter_rows(values_only=True):
        line = " ".join(str(v) for v in row if v is not None)
        if re.search(r"nonenal|methanethiol|stale", line, re.I):
            found += 1
            if found <= 3:
                print(f"  [{sh}] {line[:150]}")
    print(f"  [{sh}] 命中 {found} 行")

print("\n== 3) 精选 PDF（26 份讲义 + 7 篇论文） ==")
pdfs = list((ROOT / "资料/原始资料/精选").rglob("*.pdf"))
print("  PDF 数:", len(pdfs))
for p in pdfs:
    try:
        doc = pymupdf.open(p)
        txt = "".join(pg.get_text() for pg in doc)
    except Exception as e:
        print("  读取失败", p, e)
        continue
    ms = [m.group(0) for m in KEY.finditer(txt)]
    if ms:
        print(f"  {p.name}: {len(ms)} 次命中，示例 {set(ms[:5])}")
