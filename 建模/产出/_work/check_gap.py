# -*- coding: utf-8 -*-
# 本程序及代码是在人工智能工具辅助下完成的。
# AI 工具名称：DeepSeek Harness；模型：deepseek-flash；开发机构/公司：DeepSeek。
"""复核缺口三项在 8.xlsx 中的定位（行号以 sheet 第 1 行为表头计）。"""
import openpyxl

wb = openpyxl.load_workbook("数据/A/8.xlsx", data_only=True)
ws = wb["Sensory thresholds"]
for r in range(2, ws.max_row + 1):
    py = str(ws.cell(r, 1).value or "")
    if py.lower() in ("(e)-2-nonenal", "mesh", "stale hops") or "stale" in py.lower() or "nonenal" in py.lower():
        print(r, [ws.cell(r, c).value for c in range(1, 9)], "| mean=", ws.cell(r, 10).value, "med=", ws.cell(r, 11).value)
print("max_row", ws.max_row)

# 全表搜 stale
hits = []
for r in range(2, ws.max_row + 1):
    row = " ".join(str(ws.cell(r, c).value) for c in range(1, 5) if ws.cell(r, c).value)
    if "stale" in row.lower():
        hits.append((r, row))
print("stale hits:", hits)
