# -*- coding: utf-8 -*-
# 本程序及代码是在人工智能工具辅助下完成的。
# AI 工具名称：DeepSeek Harness；模型：deepseek-flash；开发机构/公司：DeepSeek。
"""核对若干可疑映射在 8.xlsx 中的原始行。"""
import openpyxl

wb = openpyxl.load_workbook("数据/A/8.xlsx", data_only=True)
ws = wb["Sensory thresholds"]
hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
targets = {"dihydroisocaryophyllene epoxide", "selin-11-en-4-alpha-ol", "2-ethylhexyl salicylate",
           "delta-terpinyl acetate", "ethyl 3-hexenoate", "ethyl 2-hexenoate", "Tobacarol", "Kongol"}
print(hdr)
for r in range(2, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    text = " ".join(str(v) for v in vals if v)
    if any(t.lower() in text.lower() for t in ("caryophyllene", "selin", "salicylate", "terpinyl",
                                               "hexenoate", "hexanoate", "kongol", "tobacarol")):
        print(vals[:8])
